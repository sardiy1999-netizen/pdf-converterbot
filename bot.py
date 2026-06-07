import os
import logging
import tempfile
import requests
from datetime import datetime
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Application, CommandHandler, MessageHandler,
    CallbackQueryHandler, ContextTypes, filters, ConversationHandler
)

BOT_TOKEN = "8694013926:AAEIemYge9Vrw1H7-E9BiDSztvf9Tvpqyk8"
LOG_CHANNEL_ID = -1003705652043
ILOVEPDF_PUBLIC_KEY = "project_public_ffa22f6c101db1ea8fc755f86b1144fb_FvKnmfb7adae08431d7858b2afd61a17f59de"
TEMP_DIR = tempfile.gettempdir()

logging.basicConfig(format='%(asctime)s - %(levelname)s - %(message)s', level=logging.INFO)
logger = logging.getLogger(__name__)

WAITING_FILE = 1
WAITING_SECOND_FILE = 2
WAITING_PDF_OFFICE_TYPE = 3
WAITING_PASSWORD = 4
WAITING_WATERMARK_TEXT = 5
user_sessions = {}

MAIN_MENU = [
    [InlineKeyboardButton("PDF Compress",  callback_data="compress"),
     InlineKeyboardButton("PDF Merge",      callback_data="merge")],
    [InlineKeyboardButton("PDF Split",      callback_data="split"),
     InlineKeyboardButton("PDF to JPG",    callback_data="pdf2jpg")],
    [InlineKeyboardButton("JPG to PDF",    callback_data="jpg2pdf"),
     InlineKeyboardButton("Office to PDF", callback_data="office2pdf")],
    [InlineKeyboardButton("PDF to Office", callback_data="pdf2office"),
     InlineKeyboardButton("PDF to Code",   callback_data="protect")],
    [InlineKeyboardButton("PDF No Code",   callback_data="unlock"),
     InlineKeyboardButton("PDF Rotate",    callback_data="rotate")],
    [InlineKeyboardButton("Page Numbers",  callback_data="pagenumber"),
     InlineKeyboardButton("Watermark",     callback_data="watermark")],
    [InlineKeyboardButton("PDF Repair",    callback_data="repair"),
     InlineKeyboardButton("PDF to PDF/A",  callback_data="pdfa")],
    [InlineKeyboardButton("Extract Text",  callback_data="extract")],
]

PDF2OFFICE_MENU = [
    [InlineKeyboardButton("PDF to Word (.docx)",  callback_data="pdf2word")],
    [InlineKeyboardButton("PDF to Excel (.xlsx)", callback_data="pdf2excel")],
    [InlineKeyboardButton("PDF to PPT (.pptx)",   callback_data="pdf2ppt")],
    [InlineKeyboardButton("Orqaga",               callback_data="back")],
]

TOOL_DESCRIPTIONS = {
    "compress":   "PDF faylni yuboring - hajmini kichraytiraman.",
    "merge":      "PDF fayllarni yuboring (2 tadan 5 tagacha). Bir nechta faylni birdaniga yuborishingiz yoki ketma-ket yuborishingiz mumkin.",
    "split":      "Bolish uchun PDF faylni yuboring. Har bir sahifa alohida PDF bolib chiqadi.",
    "pdf2jpg":    "JPG ga otkazish uchun PDF yuboring.",
    "jpg2pdf":    "PDF ga otkazish uchun JPG yoki PNG rasm yuboring.",
    "office2pdf": "Word (.doc, .docx), Excel (.xls, .xlsx) yoki PowerPoint (.ppt, .pptx) faylni yuboring - PDF ga otkazaman.",
    "pdf2word":   "Word (.docx) ga otkazish uchun PDF yuboring.",
    "pdf2excel":  "Excel (.xlsx) ga otkazish uchun PDF yuboring.",
    "pdf2ppt":    "PowerPoint (.pptx) ga otkazish uchun PDF yuboring.",
    "protect":    "PDF ga parol qoyish uchun faylni yuboring. Keyin o'zingiz parol kiritasiz.",
    "unlock":     "PDF dan parolni olib tashlash uchun faylni yuboring.",
    "rotate":     "Aylantirish uchun PDF yuboring (90 daraja).",
    "pagenumber": "Sahifa raqami qoshish uchun PDF yuboring.",
    "watermark":  "Watermark qoshish uchun PDF yuboring, keyin watermark matnini kiritasiz.",
    "repair":     "Buzilgan PDF faylni yuboring.",
    "pdfa":       "PDF/A formatiga otkazish uchun PDF yuboring.",
    "extract":    "Matn chiqarish uchun PDF yuboring.",
}

TOOL_DISPLAY = {
    "compress":   "PDF COMPRESS",
    "merge":      "PDF MERGE",
    "pdf2jpg":    "PDF TO JPG",
    "jpg2pdf":    "JPG TO PDF",
    "office2pdf": "OFFICE TO PDF",
    "pdf2word":   "PDF TO WORD",
    "pdf2excel":  "PDF TO EXCEL",
    "pdf2ppt":    "PDF TO PPT",
    "protect":    "PDF TO CODE",
    "unlock":     "PDF NO CODE",
    "rotate":     "PDF ROTATE",
    "pagenumber": "PAGE NUMBERS",
    "watermark":  "WATERMARK",
    "repair":     "PDF REPAIR",
    "pdfa":       "PDF TO PDF/A",
    "extract":    "EXTRACT TEXT",
}

ILOVEPDF_TOOLS = {
    "compress":   "compress",
    "merge":      "merge",
    "split":      "split",
    "pdf2jpg":    "pdfjpg",
    "jpg2pdf":    "imagepdf",
    "office2pdf": "officepdf",
    "protect":    "protect",
    "unlock":     "unlock",
    "rotate":     "rotate",
    "pagenumber": "pagenumber",
    "watermark":  "watermark",
    "repair":     "repair",
    "pdfa":       "pdfa",
    "extract":    "extract",
}

LOCAL_TOOLS = {"pdf2word", "pdf2excel", "pdf2ppt"}

OUTPUT_EXT = {
    "pdf2jpg":    "jpg",
    "pdf2word":   "docx",
    "pdf2excel":  "xlsx",
    "pdf2ppt":    "pptx",
    "jpg2pdf":    "pdf",
    "office2pdf": "pdf",
    "split":      "zip",
    "extract":    "txt",
}

ALLOWED_EXT_OFFICE = [".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx"]


async def send_action_log(context, user, action_text: str):
    now  = datetime.now().strftime("%d.%m.%Y")
    time = datetime.now().strftime("%H:%M:%S")
    username = f"@{user.username}" if user.username else "mavjud emas"
    full_name = user.full_name or "Noma'lum"

    message = f"""
<b>⚡️ FOYDALANUVCHI FAOLIYATI</b> ({now})

<b>👤 FOYDALANUVCHI:</b>
ID: <code>{user.id}</code>
Ism: <b>{full_name}</b>
Username: {username}

<b>⚙️ HARAKAT:</b>
{action_text}

<b>🕐 VAQT:</b> {time}
"""
    try:
        await context.bot.send_message(
            chat_id=LOG_CHANNEL_ID,
            text=message,
            parse_mode="HTML"
        )
    except Exception as e:
        logger.error(f"❌ Harakat logini yuborishda xatolik: {e}")


async def send_log(context, user, tool: str, file_name: str, file_path: str = None, is_photo: bool = False):
    now  = datetime.now().strftime("%d.%m.%Y")
    time = datetime.now().strftime("%H:%M:%S")
    username = f"@{user.username}" if user.username else "mavjud emas"
    full_name = user.full_name or "Noma'lum"
    
    caption_text = f"""
<b>📢 YANGI FAYL YUKLANDI</b> ({now})

<b>👤 FOYDALANUVCHI:</b>
ID: <code>{user.id}</code>
Ism: <b>{full_name}</b>
Username: {username}

<b>⚙️ AMAL:</b>
Funksiya: <b>{TOOL_DISPLAY.get(tool, tool.upper())}</b>
Fayl: <code>{file_name}</code>

<b>🕐 VAQT:</b> {time}
"""
    try:
        if file_path and os.path.exists(file_path):
            with open(file_path, "rb") as f:
                if is_photo:
                    await context.bot.send_photo(
                        chat_id=LOG_CHANNEL_ID,
                        photo=f,
                        caption=caption_text,
                        parse_mode="HTML"
                    )
                else:
                    await context.bot.send_document(
                        chat_id=LOG_CHANNEL_ID,
                        document=f,
                        filename=file_name,
                        caption=caption_text,
                        parse_mode="HTML"
                    )
        else:
            await context.bot.send_message(
                chat_id=LOG_CHANNEL_ID,
                text=caption_text,
                parse_mode="HTML"
            )
    except Exception as e:
        logger.error(f"❌ Kanalga faylli log yuborishda xatolik: {e}")


def local_pdf2word(input_path: str, output_path: str):
    from pdf2docx import Converter
    cv = Converter(input_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()

def local_pdf2excel(input_path: str, output_path: str):
    import pdfplumber
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(input_path) as pdf:
        for page_num, page in enumerate(pdf.pages):
            ws = wb.create_sheet(title=f"Page{page_num + 1}")
            tables = page.extract_tables()
            if tables:
                for table in tables:
                    for row in table:
                        ws.append([cell if cell else "" for cell in row])
            else:
                text = page.extract_text()
                if text:
                    for line in text.split("\n"):
                        ws.append([line])
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(output_path)

def local_pdf2ppt(input_path: str, output_path: str):
    from pptx import Presentation
    from pptx.util import Inches
    import fitz
    doc = fitz.open(input_path)
    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    for page_num in range(len(doc)):
        page     = doc[page_num]
        pix      = page.get_pixmap(dpi=150)
        img_path = os.path.join(TEMP_DIR, f"slide_{page_num}.png")
        pix.save(img_path)
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, 0, 0, prs.slide_width, prs.slide_height)
        os.remove(img_path)
    doc.close()
    prs.save(output_path)

def process_local(tool: str, file_path: str, user_id: int) -> str:
    ext      = OUTPUT_EXT.get(tool, "pdf")
    out_path = os.path.join(TEMP_DIR, f"result_{user_id}_{tool}.{ext}")
    if tool == "pdf2word":
        local_pdf2word(file_path, out_path)
    elif tool == "pdf2excel":
        local_pdf2excel(file_path, out_path)
    elif tool == "pdf2ppt":
        local_pdf2ppt(file_path, out_path)
    return out_path

def get_auth_token():
    resp = requests.post(
        "https://api.ilovepdf.com/v1/auth",
        json={"public_key": ILOVEPDF_PUBLIC_KEY}
    )
    resp.raise_for_status()
    return resp.json()["token"]

def process_ilovepdf(tool: str, file_paths: list, file_names: list, user_id: int, extra_data: dict = None) -> str:
    token     = get_auth_token()
    task_name = ILOVEPDF_TOOLS[tool]

    r = requests.get(
        f"https://api.ilovepdf.com/v1/start/{task_name}",
        headers={"Authorization": f"Bearer {token}"}
    )
    r.raise_for_status()
    d       = r.json()
    server  = d["server"]
    task_id = d["task"]
    token   = d.get("token", token)
    headers = {"Authorization": f"Bearer {token}"}
    base    = f"https://{server}/v1"

    server_files = []
    upload_names = []
    for fp, fname in zip(file_paths, file_names):
        with open(fp, "rb") as f:
            ur = requests.post(
                f"{base}/upload",
                headers=headers,
                data={"task": task_id},
                files={"file": (fname, f, "application/octet-stream")}
            )
        ur.raise_for_status()
        server_files.append(ur.json()["server_filename"])
        upload_names.append(fname)

    files_list = []
    for sfn, fn in zip(server_files, upload_names):
        entry = {"server_filename": sfn, "filename": fn}
        if tool == "rotate":
            entry["rotate"] = 90
        files_list.append(entry)

    body = {"task": task_id, "tool": task_name, "files": files_list}

    if tool == "compress":
        body["compression_level"] = "recommended"
    elif tool == "protect" and extra_data and "password" in extra_data:
        body["password"] = extra_data["password"]
    elif tool == "pagenumber":
        body["vertical_position"] = "bottom"
        body["horizontal_position"] = "right"
        body["pages"] = "all"
    elif tool == "watermark" and extra_data and "watermark_text" in extra_data:
        body["text"] = extra_data["watermark_text"]
        body["vertical_position"] = "middle"
        body["horizontal_position"] = "center"
        body["rotation"] = 45
        body["font_size"] = 40
        body["transparency"] = 50
    elif tool == "pdfa":
        body["conformance"] = "pdfa-2b"

    pr = requests.post(f"{base}/process", headers=headers, json=body)
    pr.raise_for_status()

    dr = requests.get(f"{base}/download/{task_id}", headers=headers, stream=True)
    dr.raise_for_status()

    ext      = OUTPUT_EXT.get(tool, "pdf")
    out_path = os.path.join(TEMP_DIR, f"result_{user_id}_{tool}.{ext}")
    with open(out_path, "wb") as f:
        for chunk in dr.iter_content(chunk_size=8192):
            if chunk:
                f.write(chunk)
    return out_path

def make_result_name(orig_name: str, ext: str) -> str:
    base = os.path.splitext(orig_name)[0]
    return f"{base}_sardorpdf.{ext}"


async def show_main_menu(update: Update, context: ContextTypes.DEFAULT_TYPE, message_text: str = None):
    """Asosiy menyuni ko'rsatadigan funksiya"""
    if message_text is None:
        message_text = "🏠 Asosiy menyu:\n\nYangi amalni tanlang:"
    
    if update.callback_query:
        await update.callback_query.message.reply_text(
            message_text,
            reply_markup=InlineKeyboardMarkup(MAIN_MENU)
        )
    else:
        await update.message.reply_text(
            message_text,
            reply_markup=InlineKeyboardMarkup(MAIN_MENU)
        )

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    await send_action_log(context, user, "🤖 Botni ishga tushirdi (<b>/start</b> buyrug'ini bosdi).")

    await update.message.reply_text(
        f"Salom, {user.first_name}!\n\n"
        "Sardorpdf botiga xush kelibsiz. Ushbu bot @sardorbek_abdullaev tomonidan sizlarga manfaatli bo'lishi uchun yaratildi.\n"
        "(botda kamchilik yoki taklif uchun men bilan bog'laning)\n\n"
        "Amallardan birini tanlang:",
        reply_markup=InlineKeyboardMarkup(MAIN_MENU)
    )
    return WAITING_FILE

async def menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query   = update.callback_query
    await query.answer()
    tool    = query.data
    user_id = query.from_user.id
    user    = query.from_user

    if tool == "pdf2office":
        await send_action_log(context, user, f"📂 Submenuga kirdi: <b>PDF TO OFFICE</b>")
        await query.edit_message_text(
            "PDF ni qaysi formatga otkazmoqchisiz? Birini tanlang:",
            reply_markup=InlineKeyboardMarkup(PDF2OFFICE_MENU)
        )
        return WAITING_PDF_OFFICE_TYPE

    if tool == "merge":
        user_sessions[user_id] = {"tool": tool, "files": [], "names": [], "merge_mode": True}
        await send_action_log(context, user, f"⚙️ Amallardan birini tanladi: <b>{TOOL_DISPLAY.get(tool, tool.upper())}</b>")
        await query.edit_message_text(
            f"Tanlandi: MERGE\n\n"
            f"📌 PDF fayllarni birlashtirish uchun 2 tadan 5 tagacha PDF fayl yuboring.\n"
            f"✅ Fayllarni birdaniga yuborishingiz yoki ketma-ket yuborishingiz mumkin.\n"
            f"✅ Har bir fayl qabul qilinganda sizga nechta fayl qolgani haqida xabar beriladi.\n"
            f"✅ 5 ta fayl yuborilgandan keyin avtomatik birlashtiriladi.\n\n"
            f"⚠️ Birlashtirish uchun kamida 2 ta fayl kerak!\n\n"
            f"Fayllarni yuborishni boshlang:",
            reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("❌ Bekor qilish", callback_data="back")]])
        )
        return WAITING_FILE

    user_sessions[user_id] = {"tool": tool, "files": [], "names": []}
    
    display_name = TOOL_DISPLAY.get(tool, tool.upper())
    await send_action_log(context, user, f"⚙️ Amallardan birini tanladi: <b>{display_name}</b>")

    await query.edit_message_text(
        text=f"Tanlandi: {tool.upper()}\n\n{TOOL_DESCRIPTIONS.get(tool, 'Fayl yuboring.')}",
        reply_markup=InlineKeyboardMarkup(
            [[InlineKeyboardButton("Orqaga", callback_data="back")]]
        )
    )
    return WAITING_FILE

async def pdf2office_type_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query   = update.callback_query
    await query.answer()
    tool    = query.data
    user_id = query.from_user.id
    user    = query.from_user

    if tool == "back":
        user_sessions.pop(user_id, None)
        await send_action_log(context, user, "⬅️ Submenudan asosiy menyuga qaytdi.")
        await query.edit_message_text(
            "Asosiy menyu:",
            reply_markup=InlineKeyboardMarkup(MAIN_MENU)
        )
        return WAITING_FILE

    user_sessions[user_id] = {"tool": tool, "files": [], "names": []}
    
    display_name = TOOL_DISPLAY.get(tool, tool.upper())
    await send_action_log(context, user, f"⚙️ Office formatini tanladi: <b>{display_name}</b>")

    await query.edit_message_text(
        text=f"Tanlandi: {tool.upper()}\n\n{TOOL_DESCRIPTIONS.get(tool, 'Fayl yuboring.')}",
        reply_markup=InlineKeyboardMarkup(
            [[InlineKeyboardButton("Orqaga", callback_data="back")]]
        )
    )
    return WAITING_FILE

async def back(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    user = query.from_user
    user_sessions.pop(user.id, None)
    
    await send_action_log(context, user, "⬅️ 'Orqaga' tugmasini bosib, asosiy menyuga qaytdi.")

    await query.edit_message_text(
        "Asosiy menyu:",
        reply_markup=InlineKeyboardMarkup(MAIN_MENU)
    )
    return WAITING_FILE

async def handle_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    user    = update.effective_user
    session = user_sessions.get(user_id)

    if not session:
        await update.message.reply_text(
            "Avval amal tanlang:",
            reply_markup=InlineKeyboardMarkup(MAIN_MENU)
        )
        return WAITING_FILE

    tool = session["tool"]
    is_photo = False

    if update.message.document:
        file_obj  = update.message.document
        orig_name = file_obj.file_name or "file.pdf"
    elif update.message.photo:
        file_obj  = update.message.photo[-1]
        orig_name = "image.jpg"
        is_photo = True
    else:
        await update.message.reply_text("Fayl yoki rasm yuboring.")
        return WAITING_FILE

    if tool == "office2pdf":
        ext = os.path.splitext(orig_name)[1].lower()
        if ext not in ALLOWED_EXT_OFFICE:
            await update.message.reply_text(
                "Xato fayl turi!\n\nOffice to PDF uchun quyidagi formatlar qabul qilinadi:\n"
                "Word: .doc, .docx\n"
                "Excel: .xls, .xlsx\n"
                "PowerPoint: .ppt, .pptx"
            )
            return WAITING_FILE

    processing_msg = await update.message.reply_text("Fayl qabul qilindi, ishlanmoqda...")

    try:
        tg_file   = await file_obj.get_file()
        file_path = os.path.join(TEMP_DIR, f"{user_id}_{file_obj.file_id}.bin")
        await tg_file.download_to_drive(file_path)
        session["files"].append(file_path)
        session["names"].append(orig_name)

        await send_log(context, user, tool, orig_name, file_path, is_photo=is_photo)

        if tool == "merge" and session.get("merge_mode"):
            file_count = len(session["files"])
            if file_count < 5:
                if file_count >= 2:
                    await processing_msg.edit_text(
                        f"✅ {file_count}-fayl qabul qilindi.\n\n"
                        f"📊 Jami: {file_count} ta fayl yuborildi.\n"
                        f"⚠️ Birlashtirish uchun kamida 2 ta fayl kerak.\n"
                        f"➕ Yana {5 - file_count} ta fayl qo'shishingiz mumkin.\n"
                        f"✅ Agar yetarli deb hisoblasangiz, 'Birlashtirish' tugmasini bosing.\n\n"
                        f"Yangi fayl yuborishingiz yoki tugmani bosishingiz mumkin:",
                        reply_markup=InlineKeyboardMarkup([
                            [InlineKeyboardButton("🔗 Birlashtirish", callback_data="do_merge")],
                            [InlineKeyboardButton("❌ Bekor qilish", callback_data="back")]
                        ])
                    )
                else:
                    await processing_msg.edit_text(
                        f"✅ {file_count}-fayl qabul qilindi.\n\n"
                        f"⚠️ Hali birlashtirish uchun yetarli emas! Kamida 2 ta fayl kerak.\n"
                        f"➕ Yana {5 - file_count} ta fayl qo'shishingiz mumkin.\n\n"
                        f"Fayl yuborishda davom eting:"
                    )
                return WAITING_FILE
            else:
                await processing_msg.edit_text("5 ta fayl qabul qilindi. Birlashtirilmoqda...")
                await perform_merge(update, context, user_id, processing_msg)
                return WAITING_FILE

        if tool == "protect":
            await processing_msg.delete()
            await update.message.reply_text(
                "✅ PDF fayl qabul qilindi.\n\n"
                "🔐 PDF ga o'rnatmoqchi bo'lgan parolingizni kiriting:\n"
                "(Faqat harflar va raqamlardan iborat bo'lishi mumkin)"
            )
            return WAITING_PASSWORD

        if tool == "watermark":
            await processing_msg.delete()
            await update.message.reply_text(
                "✅ PDF fayl qabul qilindi.\n\n"
                "💧 Watermark matnini kiriting:\n"
                "(Masalan: 'SardorPDF', 'CONFIDENTIAL', 'MUHIM' va h.k.)"
            )
            return WAITING_WATERMARK_TEXT

        await processing_msg.edit_text("Ishlanmoqda, kuting...")

        if tool in LOCAL_TOOLS:
            result_path = process_local(tool, session["files"][0], user_id)
        else:
            result_path = process_ilovepdf(tool, session["files"], session["names"], user_id)

        if result_path and os.path.exists(result_path):
            ext       = OUTPUT_EXT.get(tool, "pdf")
            send_name = make_result_name(orig_name, ext)
            with open(result_path, "rb") as f:
                await update.message.reply_document(
                    document=f,
                    filename=send_name,
                    caption=f"✅ {tool.upper()} muvaffaqiyatli bajarildi!"
                )
            
            await send_action_log(context, user, f"✅ <b>{TOOL_DISPLAY.get(tool, tool.upper())}</b> jarayoni muvaffaqiyatli yakunlandi va foydalanuvchiga natija yuborildi.")
            
            # AVTOMATIK MENYU
            await show_main_menu(update, context, "✅ Fayl tayyor!\n\n🏠 Asosiy menyuga qaytdingiz. Yangi amalni tanlang:")

            try:
                os.remove(result_path)
            except:
                pass
            await processing_msg.delete()
        else:
            await processing_msg.edit_text("Natija topilmadi.")
            await show_main_menu(update, context, "❌ Xatolik yuz berdi!\n\n🏠 Asosiy menyu:")

    except requests.HTTPError as e:
        logger.error(f"HTTP xato: {e.response.text}")
        await processing_msg.edit_text(f"API xatolik: {e.response.status_code}\n{e.response.text[:300]}")
        await show_main_menu(update, context, "⚠️ API xatolik yuz berdi!\n\n🏠 Asosiy menyu:")
    except Exception as e:
        logger.error(f"Xato: {e}", exc_info=True)
        await processing_msg.edit_text(f"Xatolik: {str(e)}")
        await show_main_menu(update, context, "⚠️ Xatolik yuz berdi!\n\n🏠 Asosiy menyu:")
    finally:
        if tool not in ["protect", "watermark", "merge"]:
            for f in session.get("files", []):
                try:
                    os.remove(f)
                except:
                    pass
            user_sessions.pop(user_id, None)
    
    return WAITING_FILE

async def perform_merge(update: Update, context: ContextTypes.DEFAULT_TYPE, user_id: int, processing_msg=None):
    session = user_sessions.get(user_id)
    if not session or session["tool"] != "merge":
        return
    
    user = update.effective_user
    
    try:
        if len(session["files"]) < 2:
            await processing_msg.edit_text("❌ Birlashtirish uchun kamida 2 ta fayl kerak!")
            await show_main_menu(update, context, "❌ Kamida 2 ta fayl kerak!\n\n🏠 Asosiy menyu:")
            return
        
        result_path = process_ilovepdf("merge", session["files"], session["names"], user_id)
        
        if result_path and os.path.exists(result_path):
            send_name = make_result_name("merged", "pdf")
            with open(result_path, "rb") as f:
                if update.callback_query:
                    await update.callback_query.message.reply_document(
                        document=f,
                        filename=send_name,
                        caption=f"✅ {len(session['files'])} ta PDF fayl muvaffaqiyatli birlashtirildi!"
                    )
                else:
                    await update.message.reply_document(
                        document=f,
                        filename=send_name,
                        caption=f"✅ {len(session['files'])} ta PDF fayl muvaffaqiyatli birlashtirildi!"
                    )
            
            await send_action_log(context, user, f"✅ <b>PDF MERGE</b> ({len(session['files'])} ta fayl birlashtirildi) muvaffaqiyatli yakunlandi.")
            
            # AVTOMATIK MENYU
            await show_main_menu(update, context, "✅ Fayllar birlashtirildi!\n\n🏠 Asosiy menyuga qaytdingiz. Yangi amalni tanlang:")
            
            try:
                os.remove(result_path)
            except:
                pass
            if processing_msg:
                await processing_msg.delete()
        else:
            if processing_msg:
                await processing_msg.edit_text("❌ Birlashtirishda xatolik yuz berdi!")
            await show_main_menu(update, context, "❌ Birlashtirishda xatolik!\n\n🏠 Asosiy menyu:")
    except Exception as e:
        logger.error(f"Merge xatosi: {e}", exc_info=True)
        if processing_msg:
            await processing_msg.edit_text(f"❌ Xatolik: {str(e)}")
        await show_main_menu(update, context, f"❌ Xatolik: {str(e)[:100]}\n\n🏠 Asosiy menyu:")
    finally:
        for f in session.get("files", []):
            try:
                os.remove(f)
            except:
                pass
        user_sessions.pop(user_id, None)

async def merge_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    
    user_id = query.from_user.id
    session = user_sessions.get(user_id)
    
    if not session or session["tool"] != "merge":
        await query.edit_message_text("❌ Xatolik! Iltimos qaytadan urinib ko'ring.")
        return
    
    await query.edit_message_text("🔗 Birlashtirilmoqda, kuting...")
    await perform_merge(update, context, user_id, query.message)

async def handle_password(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    user = update.effective_user
    password = update.message.text.strip()
    session = user_sessions.get(user_id)
    
    if not session or session["tool"] != "protect":
        await update.message.reply_text(
            "❌ Xatolik! Iltimos qaytadan /start bosing va amalni tanlang."
        )
        return WAITING_FILE
    
    if len(password) < 4:
        await update.message.reply_text(
            "❌ Parol kamida 4 ta belgidan iborat bo'lishi kerak!\n\n"
            "Qaytadan parol kiriting:"
        )
        return WAITING_PASSWORD
    
    processing_msg = await update.message.reply_text("🔐 PDF ga parol o'rnatilmoqda, kuting...")
    
    try:
        result_path = process_ilovepdf("protect", session["files"], session["names"], user_id, {"password": password})
        
        if result_path and os.path.exists(result_path):
            send_name = make_result_name(session["names"][0], "pdf")
            with open(result_path, "rb") as f:
                await update.message.reply_document(
                    document=f,
                    filename=send_name,
                    caption=f"✅ PDF faylga parol muvaffaqiyatli o'rnatildi!\n🔐 Parol: `{password}`",
                    parse_mode="Markdown"
                )
            
            await send_action_log(context, user, f"✅ <b>PDF TO CODE</b> (PDFga parol o'rnatildi: {password}) muvaffaqiyatli yakunlandi.")
            
            # AVTOMATIK MENYU
            await show_main_menu(update, context, "✅ PDF faylga parol o'rnatildi!\n\n🏠 Asosiy menyuga qaytdingiz. Yangi amalni tanlang:")
            
            try:
                os.remove(result_path)
            except:
                pass
            await processing_msg.delete()
        else:
            await processing_msg.edit_text("❌ Parol o'rnatishda xatolik yuz berdi!")
            await show_main_menu(update, context, "❌ Xatolik!\n\n🏠 Asosiy menyu:")
    except Exception as e:
        logger.error(f"Protect xatosi: {e}", exc_info=True)
        await processing_msg.edit_text(f"❌ Xatolik: {str(e)}")
        await show_main_menu(update, context, f"❌ Xatolik: {str(e)[:100]}\n\n🏠 Asosiy menyu:")
    finally:
        for f in session.get("files", []):
            try:
                os.remove(f)
            except:
                pass
        user_sessions.pop(user_id, None)
    
    return WAITING_FILE

async def handle_watermark_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    user = update.effective_user
    watermark_text = update.message.text.strip()
    session = user_sessions.get(user_id)
    
    if not session or session["tool"] != "watermark":
        await update.message.reply_text(
            "❌ Xatolik! Iltimos qaytadan /start bosing va amalni tanlang."
        )
        return WAITING_FILE
    
    if len(watermark_text) < 2:
        await update.message.reply_text(
            "❌ Watermark matni kamida 2 ta belgidan iborat bo'lishi kerak!\n\n"
            "Qaytadan matn kiriting:"
        )
        return WAITING_WATERMARK_TEXT
    
    processing_msg = await update.message.reply_text("💧 Watermark qo'shilmoqda, kuting...")
    
    try:
        result_path = process_ilovepdf("watermark", session["files"], session["names"], user_id, {"watermark_text": watermark_text})
        
        if result_path and os.path.exists(result_path):
            send_name = make_result_name(session["names"][0], "pdf")
            with open(result_path, "rb") as f:
                await update.message.reply_document(
                    document=f,
                    filename=send_name,
                    caption=f"✅ Watermark muvaffaqiyatli qo'shildi!\n💧 Matn: `{watermark_text}`",
                    parse_mode="Markdown"
                )
            
            await send_action_log(context, user, f"✅ <b>WATERMARK</b> (Matn: {watermark_text}) muvaffaqiyatli qo'shildi.")
            
            # AVTOMATIK MENYU
            await show_main_menu(update, context, "✅ Watermark qo'shildi!\n\n🏠 Asosiy menyuga qaytdingiz. Yangi amalni tanlang:")
            
            try:
                os.remove(result_path)
            except:
                pass
            await processing_msg.delete()
        else:
            await processing_msg.edit_text("❌ Watermark qo'shishda xatolik yuz berdi!")
            await show_main_menu(update, context, "❌ Xatolik!\n\n🏠 Asosiy menyu:")
    except Exception as e:
        logger.error(f"Watermark xatosi: {e}", exc_info=True)
        await processing_msg.edit_text(f"❌ Xatolik: {str(e)}")
        await show_main_menu(update, context, f"❌ Xatolik: {str(e)[:100]}\n\n🏠 Asosiy menyu:")
    finally:
        for f in session.get("files", []):
            try:
                os.remove(f)
            except:
                pass
        user_sessions.pop(user_id, None)
    
    return WAITING_FILE

async def handle_second_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Eski merge uchun - hozir ishlatilmaydi, lekin saqlab qo'yilgan"""
    user_id = update.effective_user.id
    user    = update.effective_user
    session = user_sessions.get(user_id)

    if not session or session["tool"] != "merge":
        return await handle_file(update, context)

    if not update.message.document:
        await update.message.reply_text("PDF fayl yuboring.")
        return WAITING_SECOND_FILE

    file_obj       = update.message.document
    orig_name      = file_obj.file_name or "file2.pdf"
    processing_msg = await update.message.reply_text("Birlashtirilmoqda...")

    try:
        tg_file   = await file_obj.get_file()
        file_path = os.path.join(TEMP_DIR, f"{user_id}_{file_obj.file_id}.bin")
        await tg_file.download_to_drive(file_path)
        session["files"].append(file_path)
        session["names"].append(orig_name)

        await send_log(context, user, "merge (2-fayl)", orig_name, file_path, is_photo=False)

        await processing_msg.edit_text("Ishlanmoqda, kuting...")
        result_path = process_ilovepdf("merge", session["files"], session["names"], user_id)

        if result_path and os.path.exists(result_path):
            first_name = session["names"][0] if session["names"] else "file.pdf"
            send_name  = make_result_name(first_name, "pdf")
            with open(result_path, "rb") as f:
                await update.message.reply_document(
                    document=f,
                    filename=send_name,
                    caption="MERGE muvaffaqiyatli bajarildi!"
                )
            
            await send_action_log(context, user, "✅ <b>PDF MERGE</b> (fayllarni birlashtirish) muvaffaqiyatli yakunlandi.")
            
            # AVTOMATIK MENYU
            await show_main_menu(update, context, "✅ Fayllar birlashtirildi!\n\n🏠 Asosiy menyu:")

            try:
                os.remove(result_path)
            except:
                pass
            await processing_msg.delete()
        else:
            await processing_msg.edit_text("Xatolik yuz berdi.")
            await show_main_menu(update, context, "❌ Xatolik!\n\n🏠 Asosiy menyu:")

    except requests.HTTPError as e:
        logger.error(f"HTTP xato: {e.response.text}")
        await processing_msg.edit_text(f"API xatolik: {e.response.status_code}\n{e.response.text[:300]}")
        await show_main_menu(update, context, "⚠️ API xatolik!\n\n🏠 Asosiy menyu:")
    except Exception as e:
        logger.error(f"Xato: {e}", exc_info=True)
        await processing_msg.edit_text(f"Xatolik: {str(e)}")
        await show_main_menu(update, context, f"❌ Xatolik: {str(e)[:100]}\n\n🏠 Asosiy menyu:")
    finally:
        for f in session.get("files", []):
            try:
                os.remove(f)
            except:
                pass
        user_sessions.pop(user_id, None)

    return WAITING_FILE

def main():
    app = Application.builder().token(BOT_TOKEN).build()
    conv_handler = ConversationHandler(
        entry_points=[
            CommandHandler("start", start),
            CallbackQueryHandler(menu, pattern="^(?!back$|do_merge$).+"),
            CallbackQueryHandler(merge_callback, pattern="^do_merge$"),
        ],
        states={
            WAITING_FILE: [
                MessageHandler(filters.Document.ALL | filters.PHOTO, handle_file),
                CallbackQueryHandler(menu, pattern="^(?!back$|do_merge$).+"),
                CallbackQueryHandler(merge_callback, pattern="^do_merge$"),
                CallbackQueryHandler(back, pattern="^back$"),
            ],
            WAITING_SECOND_FILE: [
                MessageHandler(filters.Document.ALL, handle_second_file),
                CallbackQueryHandler(back, pattern="^back$"),
            ],
            WAITING_PDF_OFFICE_TYPE: [
                CallbackQueryHandler(
                    pdf2office_type_selected,
                    pattern="^(pdf2word|pdf2excel|pdf2ppt|back)$"
                ),
            ],
            WAITING_PASSWORD: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, handle_password),
            ],
            WAITING_WATERMARK_TEXT: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, handle_watermark_text),
            ],
        },
        fallbacks=[
            CommandHandler("start", start),
            CallbackQueryHandler(back, pattern="^back$"),
        ],
    )
    app.add_handler(conv_handler)
    print("Bot ishga tushdi...")
    print(f"Log kanali ID: {LOG_CHANNEL_ID}")
    print("✅ To'liq monitoring tizimi faollashtirildi!")
    print("✅ PDF to Code - foydalanuvchi parol kiritadi")
    print("✅ PDF Merge - 2-5 ta faylni qabul qiladi")
    print("✅ Watermark - foydalanuvchi matn kiritadi")
    print("✅ AVTOMATIK MENYU - har bir operatsiyadan keyin menyu avtomatik chiqadi")
    app.run_polling(allowed_updates=Update.ALL_TYPES)

if __name__ == "__main__":
    main()
